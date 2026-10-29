"""Write the rows to CSV and XLSX so that the same rows always give the same bytes.

CSV: semicolon separated, comma for cents, UTF-8 with BOM, so Excel in Spanish opens
it as it is. XLSX: numbers as numbers and dates as dates. Its zip is rebuilt with a
fixed timestamp and no compression, so two runs (or two machines) give the same file.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import rules
from .invoice import COLUMNS, Row

FIXED_TIME = dt.datetime(2026, 1, 1, 0, 0, 0)
ZIP_TIME = (1980, 1, 1, 0, 0, 0)
MONEY = {"base", "vat", "withholding", "total"}
RATES = {"vat_rate", "wh_rate"}


def _text(row: Row, attr: str) -> str:
    value = getattr(row, attr)
    if value is None:
        return ""
    if attr in MONEY:
        return rules.fmt(value)
    if attr in RATES:
        return str(value).replace(".", ",")
    if isinstance(value, dt.date):
        return value.strftime("%d/%m/%Y")
    return str(value)


def write_csv(rows: list[Row], path: Path) -> None:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([title for _, title in COLUMNS])
    for row in rows:
        writer.writerow([_text(row, attr) for attr, _ in COLUMNS])
    path.write_bytes(b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8"))


def _xlsx_bytes(rows: list[Row]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"
    ws.append([title for _, title in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    review = PatternFill("solid", fgColor="FCE4D6")
    for row in rows:
        ws.append([getattr(row, attr) for attr, _ in COLUMNS])
        excel_row = ws.max_row
        for col, (attr, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=excel_row, column=col)
            if attr in MONEY:
                cell.number_format = "#,##0.00"
            elif attr == "date":
                cell.number_format = "DD/MM/YYYY"
            if row.problems:
                cell.fill = review
    widths = [16, 34, 12, 12, 12, 12, 9, 12, 14, 12, 12, 10, 90]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
    ws.freeze_panes = "A2"
    wb.properties.creator = "invoice-to-sheet"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _normalize_zip(data: bytes) -> bytes:
    """Same members, same order, fixed dates, stored without compression."""
    stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ")
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = re.sub(rb"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?Z",
                                 stamp.encode(), content)
            member = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            member.external_attr = 0o644 << 16
            member.create_system = 3
            dst.writestr(member, content)
    return out.getvalue()


def write_xlsx(rows: list[Row], path: Path) -> None:
    path.write_bytes(_normalize_zip(_xlsx_bytes(rows)))
